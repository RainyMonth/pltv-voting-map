#!/usr/bin/env python3
"""Validate the PLTV workbook and build browser-ready map data."""

from __future__ import annotations

import json
import math
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "source" / "Data_Master.xlsx"
BOUNDARY_PATH = ROOT / "source" / "Political_Divisions.geojson"
OUTPUT_DIR = ROOT / "dist" / "data"

REQUIRED_COLUMNS = {
    "Division",
    "Division and Dorms",
    "Election Type",
    "Ward",
    "Year",
    "In Person Count",
    "Mail/Absentee Count",
    "Provisional Count",
    "Vote Count",
}

TARGET_DIVISIONS = {
    "2703",
    "2706",
    "2711",
    "2718",
    "2719",
    "2720",
    "2721",
    "2722",
}


def fail(message: str) -> None:
    raise ValueError(message)


def as_int(value, label: str, row_number: int) -> int:
    if value is None or value == "":
        fail(f"Row {row_number}: {label} is blank.")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Row {row_number}: {label} must be numeric, got {value!r}.") from exc
    if not math.isfinite(number) or not number.is_integer():
        fail(f"Row {row_number}: {label} must be a whole number, got {value!r}.")
    return int(number)


def split_dorms(combined, division: int) -> str:
    """Remove only the leading division label; preserve the workbook's dorm wording."""
    text = "" if combined is None else str(combined).strip()
    if not text:
        return "Not listed"
    pattern = rf"^\s*{re.escape(str(division))}\s*[-–—]\s*"
    dorms = re.sub(pattern, "", text, count=1).strip()
    return dorms or "Not listed"


def find_data_sheet(workbook):
    matches = []
    for worksheet in workbook.worksheets:
        headers = {
            str(cell.value).strip()
            for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
            if cell.value is not None
        }
        if REQUIRED_COLUMNS.issubset(headers):
            matches.append(worksheet)
    if not matches:
        fail("No worksheet contains all required Data Master columns.")
    matches.sort(key=lambda worksheet: worksheet.max_row, reverse=True)
    if len(matches) > 1 and matches[0].max_row == matches[1].max_row:
        fail(
            "More than one worksheet has the required columns and the same row count; "
            "the source data sheet cannot be identified safely."
        )
    return matches[0]


def build_voting_data() -> list[dict]:
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    worksheet = find_data_sheet(workbook)
    header_values = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    columns = {
        str(value).strip(): index
        for index, value in enumerate(header_values)
        if value is not None and str(value).strip()
    }

    records = []
    totals = {}
    component_mismatches = []

    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(value is None for value in row):
            continue

        raw_division = row[columns["Division"]]
        election_type = str(row[columns["Election Type"]] or "").strip().title()
        year = as_int(row[columns["Year"]], "Year", row_number)

        if election_type not in {"Primary", "General"}:
            fail(
                f"Row {row_number}: Election Type must be Primary or General, "
                f"got {election_type!r}."
            )

        vote_count = as_int(row[columns["Vote Count"]], "Vote Count", row_number)

        if str(raw_division).strip().lower() == "total":
            totals[(year, election_type)] = vote_count
            continue

        ward = as_int(row[columns["Ward"]], "Ward", row_number)
        division = as_int(raw_division, "Division", row_number)
        division_id = f"{ward:02d}{division:02d}"

        if division_id not in TARGET_DIVISIONS:
            fail(
                f"Row {row_number}: unexpected division {division_id}. "
                "This project is configured for Ward 27 divisions 03, 06, 11, 18–22."
            )

        components = [
            as_int(row[columns[name]], name, row_number)
            for name in ("In Person Count", "Mail/Absentee Count", "Provisional Count")
        ]
        if sum(components) != vote_count:
            component_mismatches.append(
                f"row {row_number}: components total {sum(components)}, Vote Count is {vote_count}"
            )

        records.append(
            {
                "year": year,
                "electionType": election_type,
                "ward": ward,
                "division": division,
                "divisionId": division_id,
                "voteCount": vote_count,
                "dorms": split_dorms(row[columns["Division and Dorms"]], division),
            }
        )

    if component_mismatches:
        fail("Vote Count validation failed: " + "; ".join(component_mismatches[:8]))

    keys = [(r["year"], r["electionType"], r["divisionId"]) for r in records]
    duplicates = [key for key, count in Counter(keys).items() if count > 1]
    if duplicates:
        fail(f"Duplicate Year × Election Type × Division rows: {duplicates[:8]}")

    groups = defaultdict(list)
    for record in records:
        groups[(record["year"], record["electionType"])].append(record)

    for key, group in sorted(groups.items()):
        found = {record["divisionId"] for record in group}
        if found != TARGET_DIVISIONS:
            missing = sorted(TARGET_DIVISIONS - found)
            extra = sorted(found - TARGET_DIVISIONS)
            fail(f"{key}: missing divisions {missing}; unexpected divisions {extra}.")
        expected_total = sum(record["voteCount"] for record in group)
        if key not in totals:
            fail(f"{key}: Total row is missing.")
        if totals[key] != expected_total:
            fail(f"{key}: Total row is {totals[key]}, but divisions sum to {expected_total}.")

    records.sort(key=lambda r: (r["year"], r["electionType"], r["division"]))
    return records


def build_boundary_data() -> dict:
    with BOUNDARY_PATH.open("r", encoding="utf-8") as handle:
        geojson = json.load(handle)

    if geojson.get("type") != "FeatureCollection":
        fail("Political_Divisions.geojson must be a GeoJSON FeatureCollection.")

    selected = []
    for feature in geojson.get("features", []):
        properties = feature.get("properties") or {}
        lower_properties = {str(key).lower(): value for key, value in properties.items()}
        division_id = lower_properties.get("division_num") or lower_properties.get("division_n")
        division_id = "" if division_id is None else str(division_id).zfill(4)
        if division_id in TARGET_DIVISIONS:
            selected.append(
                {
                    "type": "Feature",
                    "properties": {"divisionId": division_id},
                    "geometry": feature.get("geometry"),
                }
            )

    found = {feature["properties"]["divisionId"] for feature in selected}
    if found != TARGET_DIVISIONS:
        fail(
            "Boundary file does not contain all required divisions. "
            f"Missing: {sorted(TARGET_DIVISIONS - found)}"
        )
    if len(selected) != len(TARGET_DIVISIONS):
        fail("Boundary file contains duplicate features for at least one required division.")

    selected.sort(key=lambda feature: feature["properties"]["divisionId"])
    return {"type": "FeatureCollection", "features": selected}


def main() -> None:
    if not WORKBOOK_PATH.exists():
        fail(f"Missing workbook: {WORKBOOK_PATH}")
    if not BOUNDARY_PATH.exists():
        fail(f"Missing boundary file: {BOUNDARY_PATH}")

    records = build_voting_data()
    boundaries = build_boundary_data()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    with (OUTPUT_DIR / "voting.json").open("w", encoding="utf-8") as handle:
        json.dump(records, handle, ensure_ascii=False, indent=2)
        handle.write("\n")

    with (OUTPUT_DIR / "ward27_divisions.geojson").open("w", encoding="utf-8") as handle:
        json.dump(boundaries, handle, ensure_ascii=False, separators=(",", ":"))
        handle.write("\n")

    years = sorted({record["year"] for record in records})
    election_types = sorted({record["electionType"] for record in records})
    print(
        f"Built {len(records)} records and {len(boundaries['features'])} boundaries; "
        f"years {years[0]}–{years[-1]}; election types {', '.join(election_types)}."
    )


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"Data build failed: {error}", file=sys.stderr)
        raise SystemExit(1)
